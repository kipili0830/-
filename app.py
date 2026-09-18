import streamlit as st
import google.generativeai as genai
import json
import os
import io
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ==========================================
# 0. 디자인 토큰 (현대렌탈케어 SCM 대시보드 톤앤매너)
# ==========================================
NAVY_HEX = "1B2A4B"
RED_HEX = "E5493E"
BLUE_HEX = "2F6FED"
ORANGE_HEX = "E68A1E"
GRAY_HEX = "6B7280"
LIGHT_GRAY_HEX = "F4F6FA"

NAVY = RGBColor(0x1B, 0x2A, 0x4B)
RED = RGBColor(0xE5, 0x49, 0x3E)
BLUE = RGBColor(0x2F, 0x6F, 0xED)
ORANGE = RGBColor(0xE6, 0x8A, 0x1E)
GRAY = RGBColor(0x6B, 0x72, 0x80)

# ==========================================
# 1. 환경 설정 및 보안
# ==========================================
st.set_page_config(page_title="재무제표 AI 분석 시스템", page_icon="📊", layout="wide")

# ---- 전역 스타일 (대시보드 톤앤매너 적용) ----
# 참고: data-testid 기반 선택자는 Streamlit 버전에 따라 달라질 수 있습니다.
# 버전이 크게 다르면 아래 선택자만 최신 testid로 교체하면 됩니다.
CUSTOM_CSS = f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@400;500;700;900&display=swap');

html, body, [class*="css"] {{
    font-family: 'Noto Sans KR', -apple-system, BlinkMacSystemFont, sans-serif;
}}

.stApp {{
    background-color: #{LIGHT_GRAY_HEX};
}}

/* ---- 상단 헤더 (브랜드 바) ---- */
.top-header {{
    display: flex;
    justify-content: space-between;
    align-items: flex-start;
    background: #FFFFFF;
    border-radius: 14px;
    padding: 22px 28px;
    margin-bottom: 20px;
    box-shadow: 0 2px 10px rgba(27,42,75,0.07);
    border-bottom: 4px solid #{NAVY_HEX};
}}
.top-header-eyebrow {{
    color: #8A93A6;
    font-size: 12px;
    font-weight: 700;
    letter-spacing: 1.5px;
    text-transform: uppercase;
    margin-bottom: 6px;
}}
.top-header-title {{
    color: #{NAVY_HEX};
    font-size: 26px;
    font-weight: 800;
    margin-bottom: 6px;
    line-height: 1.3;
}}
.top-header-sub {{
    color: #{GRAY_HEX};
    font-size: 13px;
}}
.top-header-right {{
    text-align: right;
    white-space: nowrap;
}}
.badge-scm {{
    display: inline-block;
    background: #{NAVY_HEX};
    color: #FFFFFF;
    font-size: 12px;
    font-weight: 800;
    letter-spacing: 1px;
    padding: 6px 14px;
    border-radius: 6px;
}}
.top-header-right-sub {{
    color: #8A93A6;
    font-size: 11px;
    margin-top: 6px;
}}

/* ---- 섹션 헤더 (국문 | English 스타일) ---- */
.section-header {{
    display: flex;
    align-items: baseline;
    gap: 10px;
    margin: 26px 0 14px 0;
    padding-bottom: 10px;
    border-bottom: 2px solid #E3E7EF;
}}
.section-title {{
    color: #{NAVY_HEX};
    font-size: 18px;
    font-weight: 800;
}}
.section-divider {{ color: #C7CCD9; font-size: 16px; }}
.section-title-en {{
    color: #9098A8;
    font-size: 12px;
    font-weight: 600;
    letter-spacing: 0.5px;
}}
.sub-header {{
    color: #{NAVY_HEX};
    font-size: 15px;
    font-weight: 800;
    margin: 18px 0 10px 0;
}}

/* ---- KPI 카드형 지표 (st.metric) ---- */
div[data-testid="stMetric"] {{
    background: #FFFFFF;
    border-radius: 12px;
    padding: 16px 14px 14px 14px;
    box-shadow: 0 2px 8px rgba(27,42,75,0.07);
    border-top: 3px solid #{NAVY_HEX};
}}
div[data-testid="stMetricLabel"] {{ color: #{GRAY_HEX}; font-weight: 700; }}
div[data-testid="stMetricValue"] {{ color: #{NAVY_HEX}; font-weight: 800; }}

/* ---- 버튼 ---- */
.stButton > button, .stDownloadButton > button {{
    border-radius: 8px;
    font-weight: 700;
    border: none;
    padding: 0.55em 1.1em;
}}
.stButton > button[kind="primary"] {{
    background-color: #{NAVY_HEX};
}}
.stButton > button[kind="primary"]:hover {{
    background-color: #29406E;
}}
.stDownloadButton > button {{
    background-color: #{RED_HEX};
    color: #FFFFFF;
}}
.stDownloadButton > button:hover {{
    background-color: #C93A30;
    color: #FFFFFF;
}}

/* ---- 카드형 컨테이너(필터 패널 느낌) ---- */
div[data-testid="stVerticalBlockBorderWrapper"] {{
    background: #FAFBFD;
    border-radius: 12px;
}}

/* ---- 알림 박스 ---- */
div[data-testid="stAlert"] {{
    border-radius: 10px;
}}

/* ---- 사이드바 ---- */
section[data-testid="stSidebar"] {{
    background-color: #{NAVY_HEX};
}}
section[data-testid="stSidebar"] * {{
    color: #FFFFFF !important;
}}
section[data-testid="stSidebar"] input {{
    color: #{NAVY_HEX} !important;
    background-color: #FFFFFF !important;
    border-radius: 6px;
}}
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)


def render_top_header():
    st.markdown(
        f"""
        <div class="top-header">
            <div class="top-header-left">
                <div class="top-header-eyebrow">현대렌탈케어 SCM</div>
                <div class="top-header-title">📊 재무제표 AI 분석 &amp; 5대 지표 진단 시스템</div>
                <div class="top-header-sub">협력사 재무제표 PDF를 업로드하면 AI가 핵심 지표를 자동 산출하고 재무 건전성을 진단합니다.</div>
            </div>
            <div class="top-header-right">
                <span class="badge-scm">FINANCE AI</span>
                <div class="top-header-right-sub">Financial Diagnosis Report</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def section_header(title_kr, title_en):
    st.markdown(
        f"""
        <div class="section-header">
            <span class="section-title">{title_kr}</span>
            <span class="section-divider">|</span>
            <span class="section-title-en">{title_en}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )


def sub_header(text):
    st.markdown(f'<div class="sub-header">{text}</div>', unsafe_allow_html=True)


render_top_header()

st.sidebar.title("🔑 시스템 설정")
api_key_input = st.sidebar.text_input("Gemini API 키를 입력하세요", type="password")
API_KEY = api_key_input or st.secrets.get("GEMINI_API_KEY") or os.getenv("GEMINI_API_KEY")

if not API_KEY:
    st.info("👈 왼쪽 사이드바에 발급받은 Gemini API 키를 입력하면 시스템이 활성화됩니다.")
    st.stop()

genai.configure(api_key=API_KEY)

# ==========================================
# 2. 유틸리티 함수
# ==========================================
def safe_div(numerator, denominator):
    try:
        if denominator == 0 or denominator is None:
            return 0.0
        return float(numerator) / float(denominator)
    except (TypeError, ValueError):
        return 0.0

def extract_financial_data(pdf_bytes):
    model = genai.GenerativeModel('gemini-3.6-flash')
    prompt = """
    당신은 20년 경력의 재무 분석가입니다. 첨부된 재무제표 PDF를 읽고 다음 9가지 항목의 숫자를 추출하세요:
    1. 총자산, 2. 총부채, 3. 자기자본(자본총계), 4. 유동자산, 5. 유동부채, 6. 장단기차입금(장기차입금+단기차입금), 7. 매출액, 8. 영업이익, 9. 이자비용
    [조건]
    - 단위(예: 백만원)를 파악하고, 무조건 '1원 단위의 절대금액(정수)'으로 변환하세요.
    - 반드시 아래의 순수 JSON 형식으로만 답변하세요.
    {"총자산": 0, "총부채": 0, "자기자본": 0, "유동자산": 0, "유동부채": 0, "장단기차입금": 0, "매출액": 0, "영업이익": 0, "이자비용": 0}
    """
    response = model.generate_content([{"mime_type": "application/pdf", "data": pdf_bytes}, prompt])
    try:
        clean_text = response.text.replace("```json", "").replace("```", "").strip()
        return json.loads(clean_text)
    except:
        return None

def generate_financial_report(metrics_data):
    model = genai.GenerativeModel('gemini-3.6-flash')
    debt_ratio_str = metrics_data['부채비율'] if isinstance(metrics_data['부채비율'], str) else f"{metrics_data['부채비율']:.1f}%"
    prompt = f"""
    당신은 기업의 재무 건전성을 평가하는 시니어 재무 분석가입니다.
    다음 산출된 5대 재무 지표를 심층 진단하세요.
    [산출 데이터]
    - 부채비율: {debt_ratio_str}
    - 유동비율: {metrics_data['유동비율']:.1f}%
    - 차입금의존도: {metrics_data['차입금의존도']:.1f}%
    - 매출액영업이익률: {metrics_data['매출액영업이익률']:.1f}%
    - 이자보상배율: {metrics_data['이자보상배율']:.1f}배

    [작성 규칙]
    1. 각 지표별로 "양호", "보통", "주의", "위험", "적자", "부도 위험", "자본잠식" 중 가장 적합한 키워드를 골라 반드시 `- [키워드]` 형태로 요약 결론을 내려주세요.
    2. 긍정적 내용은 `:blue[파란색]`, 부정적 내용은 `:red[빨간색]`, 중간은 `:orange[주황색]` Streamlit 태그로 시각적 강조를 해주세요.

    [작성 예시 - 반드시 이 형식을 지키세요!]
    (1) 부채비율: 292.5% - :red[위험]
    - 기준: 일반 제조업 권장 기준 200% 이하
    - 진단: 권장 기준을 크게 초과하여 타인자본 의존도가 매우 높은 :red[불안정한 상태]입니다.
    (2) 유동비율: 109.2% - :orange[주의]
    - 기준: 일반 제조업 권장 기준 150% 이상
    - 진단: 100%를 겨우 넘어 단기 채무를 가까스로 상환할 수 있는 :orange[아슬아슬한 유동성 수준]입니다.

    한국통계시스템의 최신 중견기업 제조업 평균 통계를 모를 경우 "일반적인 제조업 재무 건전성 기준"으로 평가한다고 명시하세요.
    """
    response = model.generate_content(prompt)
    return response.text

def extract_status_from_report(report_text, metric_name):
    pattern = re.compile(f".*?{metric_name}.*?(-|:).*?(:red|:blue|:orange)\[(.*?)\]", re.DOTALL)
    match = pattern.search(report_text)
    if match:
        status_text = match.group(3)
        if "위험" in status_text or "적자" in status_text or "부도" in status_text or "주의" in status_text:
             return f"🚨 {status_text}"
        elif "양호" in status_text or "안전" in status_text:
             return f"🔵 {status_text}"
        else:
             return f"🟠 {status_text}"
    return "⚪️ 분석중"

# ---- 워드 문서용 스타일 헬퍼 ----
def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_horizontal_line(paragraph, color=NAVY_HEX, size="18"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

COLOR_MAP = {"red": RED, "blue": BLUE, "orange": ORANGE}
TAG_PATTERN = re.compile(r':(red|blue|orange)\[(.*?)\]')

def add_rich_text(paragraph, text, bold=False, base_size=10.5):
    """':red[...]', ':blue[...]', ':orange[...]' 태그를 워드 색상 텍스트로 변환하며 추가"""
    pos = 0
    has_content = False
    for m in TAG_PATTERN.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            run.font.size = Pt(base_size)
            run.bold = bold
            has_content = True
        color_name, content = m.group(1), m.group(2)
        run = paragraph.add_run(content)
        run.font.size = Pt(base_size)
        run.font.color.rgb = COLOR_MAP.get(color_name, GRAY)
        run.bold = True
        has_content = True
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.size = Pt(base_size)
        run.bold = bold
        has_content = True
    if not has_content:
        run = paragraph.add_run(text)
        run.font.size = Pt(base_size)
        run.bold = bold

def generate_word_document(metrics, report_text):
    """워드 문서 생성 함수 (현대렌탈케어 SCM 대시보드 톤앤매너 적용)"""
    doc = Document()

    base_style = doc.styles['Normal']
    base_style.font.name = '맑은 고딕'
    base_style.font.size = Pt(10.5)
    base_style.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    # ---- 헤더 블록 ----
    eyebrow_p = doc.add_paragraph()
    run = eyebrow_p.add_run("현대렌탈케어 SCM")
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY
    run.bold = True

    title_p = doc.add_paragraph()
    run = title_p.add_run("📊 재무제표 AI 분석 & 심층 진단 리포트")
    run.font.size = Pt(22)
    run.font.color.rgb = NAVY
    run.bold = True

    sub_p = doc.add_paragraph()
    run = sub_p.add_run("AI 기반 5대 핵심 재무 지표 자동 산출 및 건전성 진단 결과")
    run.font.size = Pt(10.5)
    run.font.color.rgb = GRAY
    run.italic = True

    date_p = doc.add_paragraph()
    run = date_p.add_run(f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    add_horizontal_line(date_p)
    date_p.paragraph_format.space_after = Pt(18)

    # ---- 1. 핵심 재무 지표 요약 (표) ----
    h1 = doc.add_paragraph()
    r = h1.add_run("1.  핵심 재무 지표 요약")
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY
    r.bold = True
    h1.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    headers = ["지표", "산출값", "진단"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(hdr_cells[i], NAVY_HEX)

    debt_str = metrics['부채비율'] if isinstance(metrics['부채비율'], str) else f"{metrics['부채비율']:.1f}%"
    rows_data = [
        ("부채비율", debt_str, extract_status_from_report(report_text, "부채비율")),
        ("유동비율", f"{metrics['유동비율']:.1f}%", extract_status_from_report(report_text, "유동비율")),
        ("차입금의존도", f"{metrics['차입금의존도']:.1f}%", extract_status_from_report(report_text, "차입금의존도")),
        ("매출액영업이익률", f"{metrics['매출액영업이익률']:.1f}%", extract_status_from_report(report_text, "매출액영업이익률")),
        ("이자보상배율", f"{metrics['이자보상배율']:.2f}배", extract_status_from_report(report_text, "이자보상배율")),
    ]
    for idx, (name, val, status) in enumerate(rows_data):
        row_cells = table.add_row().cells

        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(name)
        r0.font.size = Pt(10.5)
        r0.bold = True
        r0.font.color.rgb = NAVY

        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(val)
        r1.font.size = Pt(10.5)

        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        status_clean = status.replace("🚨", "").replace("🔵", "").replace("🟠", "").replace("⚪️", "").strip()
        r2 = p2.add_run(status_clean)
        r2.bold = True
        r2.font.size = Pt(10.5)
        if any(k in status_clean for k in ["위험", "적자", "부도", "주의", "자본잠식"]):
            r2.font.color.rgb = RED
        elif any(k in status_clean for k in ["양호", "안전"]):
            r2.font.color.rgb = BLUE
        elif "분석중" in status_clean:
            r2.font.color.rgb = GRAY
        else:
            r2.font.color.rgb = ORANGE

        if idx % 2 == 1:
            for c in row_cells:
                set_cell_background(c, LIGHT_GRAY_HEX if False else "F4F6FA")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)

    # ---- 2. AI 심층 진단 상세 내역 ----
    h2 = doc.add_paragraph()
    r = h2.add_run("2.  AI 심층 진단 상세 내역")
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY
    r.bold = True
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(10)

    for line in report_text.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        is_metric_line = bool(re.match(r'^\(\d+\)', line))
        add_rich_text(p, line, bold=is_metric_line, base_size=11 if is_metric_line else 10.5)
        p.paragraph_format.space_after = Pt(4)

    # ---- 푸터 ----
    doc.add_paragraph()
    footer_line = doc.add_paragraph()
    add_horizontal_line(footer_line, color="D9DCE3", size="6")

    note_p = doc.add_paragraph()
    r = note_p.add_run("⚠ 본 리포트는 AI가 자동 생성한 참고 자료입니다. 최종 의사결정 시 반드시 전문가 검토를 거치시기 바랍니다.")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY
    r.italic = True

    brand_p = doc.add_paragraph()
    r = brand_p.add_run("현대렌탈케어 SCM  ·  재무제표 AI 분석 시스템")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY
    r.bold = True

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# ==========================================
# 3. UI 및 워크플로우 구성
# ==========================================
fields = ["총자산", "총부채", "자기자본", "유동자산", "유동부채", "장단기차입금", "매출액", "영업이익", "이자비용"]
for field in fields:
    if field not in st.session_state:
        st.session_state[field] = 0

st.warning("⚠️ 주의: 본 시스템에는 DART(전자공시시스템) 등에 공시된 공개용 재무제표 PDF만 업로드해 주시기 바랍니다.")

section_header("1. 재무제표 PDF 업로드", "PDF Upload")
with st.container(border=True):
    uploaded_file = st.file_uploader("협력사 재무제표 PDF 파일을 업로드하세요", type=["pdf"])
    if uploaded_file is not None:
        if st.button("📄 AI로 숫자 자동 추출하기", type="primary"):
            with st.spinner("AI가 문서를 분석 중입니다..."):
                extracted_data = extract_financial_data(uploaded_file.read())
                if extracted_data:
                    for field in fields:
                        st.session_state[field] = extracted_data.get(field, 0)
                    st.success("추출 완료! 아래에서 숫자를 확인하세요.")

section_header("2. 재무 데이터 검증 및 수정", "Data Verification · 단위: 원")
with st.container(border=True):
    cols = st.columns(3)
    for idx, field in enumerate(fields):
        col = cols[idx % 3]
        st.session_state[field] = col.number_input(
            f"{field}",
            value=int(st.session_state[field]),
            step=1000000,
            format="%d"
        )
        col.caption(f"💸 **{st.session_state[field]:,}** 원")

section_header("3. 5대 재무 지표 심층 분석", "Deep-dive Diagnosis")
if st.button("🚀 5대 지표 심층 분석 실행", type="primary"):
    with st.spinner("AI가 재무 리포트를 작성 중입니다..."):
        ta, tl, eq = st.session_state["총자산"], st.session_state["총부채"], st.session_state["자기자본"]
        ca, cl, borr = st.session_state["유동자산"], st.session_state["유동부채"], st.session_state["장단기차입금"]
        rev, op, ie = st.session_state["매출액"], st.session_state["영업이익"], st.session_state["이자비용"]

        debt_ratio_val = "자본잠식" if eq < 0 else safe_div(tl, eq) * 100
        metrics = {
            "부채비율": debt_ratio_val,
            "유동비율": safe_div(ca, cl) * 100,
            "차입금의존도": safe_div(borr, ta) * 100,
            "매출액영업이익률": safe_div(op, rev) * 100,
            "이자보상배율": safe_div(op, ie)
        }

        report = generate_financial_report(metrics)
        st.session_state['report'] = report

        sub_header("🧮 핵심 지표 계산 결과 (AI 진단 연동)")
        c1, c2, c3, c4, c5 = st.columns(5)

        c1.metric("부채비율", debt_ratio_val if isinstance(debt_ratio_val, str) else f"{debt_ratio_val:.1f}%", extract_status_from_report(report, "부채비율"), delta_color="off")
        c2.metric("유동비율", f"{metrics['유동비율']:.1f}%", extract_status_from_report(report, "유동비율"), delta_color="off")
        c3.metric("차입금의존도", f"{metrics['차입금의존도']:.1f}%", extract_status_from_report(report, "차입금의존도"), delta_color="off")
        c4.metric("영업이익률", f"{metrics['매출액영업이익률']:.1f}%", extract_status_from_report(report, "매출액영업이익률"), delta_color="off")
        c5.metric("이자보상배율", f"{metrics['이자보상배율']:.2f}배", extract_status_from_report(report, "이자보상배율"), delta_color="off")

        sub_header("🤖 AI 심층 진단 리포트 (색상 시각화 적용)")
        with st.container(border=True):
            st.write(report)

        word_data = generate_word_document(metrics, report)

        st.download_button(
            label="📄 분석 결과 워드(Word) 리포트 다운로드",
            data=word_data,
            file_name="재무분석_심층리포트.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
